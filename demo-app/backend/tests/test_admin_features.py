from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, update

from app.database import SessionLocal
from app.main import app
from app.matching import normalize_text
from app.models import HistoryQuote, QuoteOption


def workbook_bytes(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "询价单"
    sheet.append(["序号", "产品名称", "参数", "型号", "单位", "数量"])
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def single_line_workbook() -> bytes:
    return workbook_bytes([[1, "电子天平", "100g，0.001g，带防风罩", "", "台", 2]])


def login(client: TestClient, username: str = "admin", password: str = "admin123"):
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text


def create_job(client: TestClient, customer_id: int, content: bytes | None = None, option_count: int = 3) -> str:
    response = client.post(
        "/api/quote-jobs",
        data={"customer_id": customer_id, "requested_option_count": option_count},
        files={"file": ("询价.xlsx", content or single_line_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 202, response.text
    return response.json()["id"]


def ordinary_customer_id(client: TestClient) -> int:
    customers = client.get("/api/customers").json()
    return next(item for item in customers if item["customer_type"] == "ordinary")["id"]


def test_export_without_confirmation_allowed():
    """未确认任何行也能导出（客户版未确认行留空），不再 409 阻断。"""
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, ordinary_customer_id(client))
        export = client.post(f"/api/quote-jobs/{job_id}/export/customer")
        assert export.status_code == 200, export.text
        workbook = load_workbook(BytesIO(export.content))
        assert "询价单" in workbook.sheetnames


def test_auto_confirm_and_export_flow():
    """一键导出：自动带价行被确认并带价导出；无自动带价行保持未确认。"""
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, ordinary_customer_id(client))
        lines = client.get(f"/api/quote-jobs/{job_id}/lines").json()
        assert lines["total"] >= 1
        auto_priced = sum(1 for it in lines["items"] if (it.get("selected_option_count") or 0) > 0)
        response = client.post(f"/api/quote-jobs/{job_id}/auto-confirm-and-export/customer")
        assert response.status_code == 200, response.text
        job = client.get(f"/api/quote-jobs/{job_id}").json()
        assert job["confirmed_lines"] == auto_priced
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook["询价单"]
        # 至少一行有确认单价（自动带价行）
        priced = [sheet.cell(2, c).value for c in range(1, sheet.max_column + 1) if isinstance(sheet.cell(2, c).value, (int, float))]
        assert any(isinstance(v, (int, float)) for v in priced)


def test_lines_sort_score_asc():
    with TestClient(app) as client:
        login(client)
        content = workbook_bytes([
            [1, "电子天平", "100g，0.001g，带防风罩", "", "台", 2],
            [2, "不存在的仪器XYZ", "定制参数abc", "", "台", 1],
        ])
        job_id = create_job(client, ordinary_customer_id(client), content)
        default = client.get(f"/api/quote-jobs/{job_id}/lines").json()
        assert default["total"] == 2
        assert default["items"][0]["source_row"] == 2
        sorted_result = client.get(f"/api/quote-jobs/{job_id}/lines?sort=score_asc").json()
        scores = [item["recommended_score"] for item in sorted_result["items"]]
        assert scores == sorted(scores)
        assert sorted_result["items"][0]["recommended_score"] <= default["items"][0]["recommended_score"]


def test_export_multi_option_quantity_and_totals():
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, ordinary_customer_id(client), option_count=3)
        lines = client.get(f"/api/quote-jobs/{job_id}/lines").json()
        line = client.get(f"/api/quote-lines/{lines['items'][0]['id']}").json()
        # select up to 3 options with distinct manufacturers
        seen = set()
        chosen = []
        for option in line["options"]:
            key = normalize_text(option["record"]["manufacturer"] or option["record"]["brand"]) or option["id"]
            if key in seen:
                continue
            seen.add(key)
            chosen.append(option["id"])
            if len(chosen) == 3:
                break
        update_response = client.patch(
            f"/api/quote-lines/{line['id']}",
            json={"selected_option_ids": chosen, "final_prices": {}, "manual_note": ""},
        )
        assert update_response.status_code == 200, update_response.text
        confirm = client.post(f"/api/quote-lines/{line['id']}/confirm", json={"override_reason": "测试"})
        assert confirm.status_code == 200, confirm.text

        export = client.post(f"/api/quote-jobs/{job_id}/export/customer")
        assert export.status_code == 200
        workbook = load_workbook(BytesIO(export.content))
        sheet = workbook["询价单"]
        header_values = [cell.value for cell in sheet[1]]
        for label in ("数量", "确认单价", "总价", "品牌", "制造商", "型号",
                      "报价2单价", "报价2制造商", "报价3单价", "报价3制造商"):
            assert label in header_values
        # 复用原表已有的 单位/数量 列，不追加重复列（旧版会追加重 数量/单位）
        normalized = [str(v).strip() for v in header_values if v]
        duplicated = [x for x in set(normalized) if normalized.count(x) > 1]
        assert not duplicated, f"导出表头出现重复列: {duplicated}"
        # 数量 / 确认单价 / 总价 written for the confirmed line (row 2)
        quantity_col = header_values.index("数量") + 1
        price_col = header_values.index("确认单价") + 1
        total_col = header_values.index("总价") + 1
        quantity = sheet.cell(2, quantity_col).value
        price = sheet.cell(2, price_col).value
        total = sheet.cell(2, total_col).value
        assert quantity == 2
        assert price and total == round(price * quantity, 2)
        # 税后单价 / 税后总价（默认税率 10%）随导出一起生成
        assert "税后单价" in header_values and "税后总价" in header_values
        tax_unit_col = header_values.index("税后单价") + 1
        tax_total_col = header_values.index("税后总价") + 1
        tax_unit = sheet.cell(2, tax_unit_col).value
        tax_total = sheet.cell(2, tax_total_col).value
        assert tax_unit == round(price * 1.1, 2)
        assert tax_total == round(price * 1.1 * quantity, 2)
        # 产品编码列存在（本测试无编码记录，应为空字符串即可，但列必须出现）
        assert "产品编码" in header_values

        detail = workbook["多方案报价明细"]
        detail_headers = [cell.value for cell in detail[1]]
        assert "数量" in detail_headers
        assert "小计" in detail_headers
        subtotal_col = detail_headers.index("小计") + 1
        price_col_d = detail_headers.index("报价") + 1
        first_data = detail[2]
        assert first_data[subtotal_col - 1].value == round(
            first_data[price_col_d - 1].value * 2, 2
        )

        internal = client.post(f"/api/quote-jobs/{job_id}/export/internal")
        assert internal.status_code == 200
        internal_wb = load_workbook(BytesIO(internal.content))
        internal_headers = [cell.value for cell in internal_wb["询价单"][1]]
        assert "匹配状态" in internal_headers
        assert "报价2单价" in internal_headers


def test_export_appends_without_overwriting_original_price():
    """原表已有“单价”列时：确认结果追加到新的“确认单价”列，原值保持不变；
    数量/总价/税后总价必须随导出生成（合并表头“参考数量”也要识别）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "清单"
    sheet.append(["序号", "产品名称", "参数", "型号", "单位", "参考数量", "单价"])
    sheet.append([1, "电子天平", "100g，0.001g，带防风罩", "", "台", 2, 999])
    stream = BytesIO()
    workbook.save(stream)
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, ordinary_customer_id(client), stream.getvalue())
        export = client.post(f"/api/quote-jobs/{job_id}/auto-confirm-and-export/internal")
        assert export.status_code == 200, export.text
        exported = load_workbook(BytesIO(export.content))
        sheet2 = exported["清单"]
        header_values = [cell.value for cell in sheet2[1]]
        original_col = header_values.index("单价") + 1
        # 原有内容不被覆盖
        assert sheet2.cell(2, original_col).value == 999
        # 确认结果在新列
        confirm_col = header_values.index("确认单价") + 1
        confirmed = sheet2.cell(2, confirm_col).value
        assert confirmed and confirmed != 999
        # 数量 / 总价 / 税后总价已生成
        quantity_col = header_values.index("参考数量") + 1
        total_col = header_values.index("总价") + 1
        tax_total_col = header_values.index("税后总价") + 1
        assert sheet2.cell(2, quantity_col).value == 2
        assert sheet2.cell(2, total_col).value == round(confirmed * 2, 2)
        assert sheet2.cell(2, tax_total_col).value == round(confirmed * 1.1 * 2, 2)


def test_job_rename_and_delete():
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, ordinary_customer_id(client))
        renamed = client.patch(f"/api/quote-jobs/{job_id}", json={"display_name": "重命名任务"})
        assert renamed.status_code == 200
        assert renamed.json()["display_name"] == "重命名任务"
        detail = client.get(f"/api/quote-jobs/{job_id}").json()
        assert detail["display_name"] == "重命名任务"
        listing = client.get("/api/quote-jobs").json()
        assert next(item for item in listing if item["id"] == job_id)["display_name"] == "重命名任务"

        deleted = client.delete(f"/api/quote-jobs/{job_id}")
        assert deleted.status_code == 200
        assert client.get(f"/api/quote-jobs/{job_id}").status_code == 404
        assert client.delete(f"/api/quote-jobs/{job_id}").status_code == 404


def test_customer_patch_and_delete():
    with TestClient(app) as client:
        login(client)
        created = client.post("/api/customers", json={"name": "转换测试客户", "customer_type": "ordinary"})
        assert created.status_code == 200
        customer_id = created.json()["id"]

        # ordinary -> special without requirements is rejected
        rejected = client.patch(f"/api/customers/{customer_id}", json={"customer_type": "special"})
        assert rejected.status_code == 400

        # ordinary -> special with requirements
        updated = client.patch(
            f"/api/customers/{customer_id}",
            json={
                "customer_type": "special",
                "requirements": [{"attribute_name": "材质", "operator": "contains", "value": "不锈钢"}],
            },
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["customer_type"] == "special"
        assert len(updated.json()["requirements"]) == 1

        # special -> vip keeps the requirements (VIP may carry requirements)
        vip = client.patch(f"/api/customers/{customer_id}", json={"customer_type": "vip"})
        assert vip.status_code == 200
        assert vip.json()["customer_type"] == "vip"
        assert len(vip.json()["requirements"]) == 1

        # requirements array is replaced wholesale
        replaced = client.patch(
            f"/api/customers/{customer_id}",
            json={"requirements": [
                {"attribute_name": "量程", "operator": ">=", "value": "100", "unit": "g"},
                {"attribute_name": "校准", "operator": "contains", "value": "外校", "required": False},
            ]},
        )
        assert replaced.status_code == 200
        assert [item["attribute_name"] for item in replaced.json()["requirements"]] == ["量程", "校准"]

        # duplicate name rejected
        other = client.post("/api/customers", json={"name": "占用名称客户"})
        assert other.status_code == 200
        conflict = client.patch(f"/api/customers/{other.json()['id']}", json={"name": "转换测试客户"})
        assert conflict.status_code == 409

        # soft delete
        deleted = client.delete(f"/api/customers/{customer_id}")
        assert deleted.status_code == 200
        names = [item["name"] for item in client.get("/api/customers").json()]
        assert "转换测试客户" not in names
        assert client.delete(f"/api/customers/{customer_id}").status_code == 404


def test_vip_customer_can_carry_requirements_and_they_apply():
    with TestClient(app) as client:
        login(client)
        created = client.post(
            "/api/customers",
            json={
                "name": "VIP带要求客户",
                "customer_type": "vip",
                "requirements": [
                    {"attribute_name": "定制属性", "operator": "contains", "value": "不可能存在的值ZZZ"}
                ],
            },
        )
        assert created.status_code == 200, created.text
        assert len(created.json()["requirements"]) == 1
        job_id = create_job(client, created.json()["id"])
        lines = client.get(f"/api/quote-jobs/{job_id}/lines").json()
        detail = client.get(f"/api/quote-lines/{lines['items'][0]['id']}").json()
        assert any(
            "未满足客户要求" in warning
            for option in detail["options"]
            for warning in option["warnings"]
        )


def test_dedup_history_redirects_option_references():
    with TestClient(app) as client:
        login(client)
        job_id = create_job(client, ordinary_customer_id(client))
        lines = client.get(f"/api/quote-jobs/{job_id}/lines").json()
        detail = client.get(f"/api/quote-lines/{lines['items'][0]['id']}").json()
        option_id = detail["options"][0]["id"]

        db = SessionLocal()
        try:
            keep = HistoryQuote(
                id="manual:dedup-keep",
                source_file="手工录入",
                name="去重测试仪器",
                normalized_name=normalize_text("去重测试仪器"),
                spec="精度0.1g",
                normalized_spec=normalize_text("精度0.1g"),
                manufacturer="测试厂",
                unit="台",
                normalized_unit=normalize_text("台"),
                price=100.0,
            )
            dup = HistoryQuote(
                id="manual:dedup-dup",
                source_file="手工录入",
                name="去重测试仪器（ ）",
                normalized_name=normalize_text("去重测试仪器（ ）"),
                spec="精度0.1g",
                normalized_spec=normalize_text("精度0.1g"),
                manufacturer="测试厂",
                unit="台",
                normalized_unit=normalize_text("台"),
                price=100.0,
            )
            db.add_all([keep, dup])
            db.execute(update(QuoteOption).where(QuoteOption.id == option_id).values(history_quote_id=dup.id))
            db.commit()
        finally:
            db.close()

        response = client.post("/api/governance/dedup-history")
        assert response.status_code == 200, response.text
        assert response.json()["removed"] == 1

        db = SessionLocal()
        try:
            assert db.get(HistoryQuote, "manual:dedup-dup") is None
            assert db.get(HistoryQuote, "manual:dedup-keep") is not None
            option = db.get(QuoteOption, option_id)
            assert option.history_quote_id == "manual:dedup-keep"
        finally:
            db.close()

        # second run is a no-op
        again = client.post("/api/governance/dedup-history")
        assert again.json()["removed"] == 0


def test_dedup_history_requires_admin():
    with TestClient(app) as client:
        login(client, "quote", "quote123")
        assert client.post("/api/governance/dedup-history").status_code == 403


def test_change_password():
    with TestClient(app) as client:
        login(client)
        created = client.post(
            "/api/users",
            json={"username": "改密测试员", "password": "initial1", "display_name": "改密测试", "role": "quote"},
        )
        assert created.status_code == 201, created.text

        with TestClient(app) as user_client:
            login(user_client, "改密测试员", "initial1")
            wrong = user_client.post(
                "/api/auth/change-password", json={"old_password": "bad", "new_password": "newpass1"}
            )
            assert wrong.status_code == 400
            short = user_client.post(
                "/api/auth/change-password", json={"old_password": "initial1", "new_password": "abc"}
            )
            assert short.status_code == 422
            ok = user_client.post(
                "/api/auth/change-password", json={"old_password": "initial1", "new_password": "newpass1"}
            )
            assert ok.status_code == 200

        with TestClient(app) as relogin:
            assert relogin.post(
                "/api/auth/login", json={"username": "改密测试员", "password": "initial1"}
            ).status_code == 401
            login(relogin, "改密测试员", "newpass1")


def test_user_management():
    with TestClient(app) as client:
        login(client)
        users = client.get("/api/users").json()
        assert any(item["username"] == "admin" for item in users)
        assert all("password_hash" not in item for item in users)

        created = client.post(
            "/api/users",
            json={"username": "管理测试员", "password": "pass1234", "display_name": "管理测试", "role": "quote"},
        )
        assert created.status_code == 201, created.text
        target = created.json()
        assert target["active"] is True

        duplicate = client.post(
            "/api/users", json={"username": "管理测试员", "password": "pass1234", "role": "quote"}
        )
        assert duplicate.status_code == 409

        # admin resets password
        reset = client.patch(f"/api/users/{target['id']}", json={"password": "reset999"})
        assert reset.status_code == 200
        with TestClient(app) as target_client:
            login(target_client, "管理测试员", "reset999")

        # deactivate blocks further API use
        deactivated = client.patch(f"/api/users/{target['id']}", json={"active": False})
        assert deactivated.status_code == 200
        assert deactivated.json()["active"] is False
        with TestClient(app) as target_client:
            login(target_client, "管理测试员", "reset999")
            assert target_client.get("/api/customers").status_code == 403

        # cannot deactivate or demote self
        me = client.get("/api/auth/me").json()
        assert client.patch(f"/api/users/{me['id']}", json={"active": False}).status_code == 400
        assert client.patch(f"/api/users/{me['id']}", json={"role": "quote"}).status_code == 400

        # role change allowed for others
        promoted = client.patch(f"/api/users/{target['id']}", json={"role": "admin", "active": True})
        assert promoted.status_code == 200
        assert promoted.json()["role"] == "admin"


def test_user_management_requires_admin():
    with TestClient(app) as client:
        login(client, "quote", "quote123")
        assert client.get("/api/users").status_code == 403
        assert client.post(
            "/api/users", json={"username": "x1", "password": "pass1234"}
        ).status_code == 403


def test_manual_history_entry():
    with TestClient(app) as client:
        login(client)
        created = client.post(
            "/api/history",
            json={
                "name": "手工录入测试仪器",
                "spec": "量程500g",
                "unit": "台",
                "manufacturer": "手工测试厂",
                "price": 888.5,
                "brand": "手工牌",
                "model": "SG-500",
                "quote_date": "2026-08",
            },
        )
        assert created.status_code == 201, created.text
        record = created.json()
        assert record["id"].startswith("manual:")
        assert record["source"] == "手工录入"

        found = client.get("/api/history/search", params={"q": "手工录入测试仪器"}).json()
        assert any(item["id"] == record["id"] for item in found)

        invalid = client.post("/api/history", json={"name": "坏记录", "price": -1})
        assert invalid.status_code == 422


def test_manual_history_entry_requires_admin():
    with TestClient(app) as client:
        login(client, "quote", "quote123")
        response = client.post("/api/history", json={"name": "越权记录", "price": 10})
        assert response.status_code == 403
