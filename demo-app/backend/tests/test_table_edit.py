"""价目表格化编辑：分页浏览接口 + 变更集保存接口。

覆盖三件事：① 表格视图要的数据能不能按页拿到；② 变更集保存的每一条
分支（成功 / 校验失败 / 版本冲突 / 去重冲突）行为是否正确；③ 改完之后
报价是不是真的用上了新数据——这是"编辑生效"的最终验收。
"""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import database as dbmod
from app import history_rows
from app.main import app
from app.models import AuditEvent, HistoryQuote
from sqlalchemy import select


def login(client: TestClient, username: str = "admin", password: str = "admin123"):
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text


def price_workbook(rows: list[tuple]) -> bytes:
    """每行 (名称, 参数, 单位, 品牌, 制造商, 单价)。

    表头刻意用「单价」而不是「含税单价」：导入路径会把含税表头的价格按固定
    税率折成税前基准价入库（excel_service.PRICE_ALIASES 逻辑），那会让断言
    里的 100 变成 90.91。本文件测的是编辑链路，不是计税。
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "价目"
    sheet.append(["序号", "产品名称", "参数", "单位", "品牌", "制造商", "单价"])
    for index, (name, spec, unit, brand, maker, price) in enumerate(rows, start=1):
        sheet.append([index, name, spec, unit, brand, maker, price])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def inquiry_workbook(name: str, spec: str, unit: str, quantity: int = 1) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "询价单"
    sheet.append(["序号", "产品名称", "参数", "型号", "单位", "数量"])
    sheet.append([1, name, spec, "", unit, quantity])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def create_database(client: TestClient, name: str) -> str:
    response = client.post("/api/databases", json={"name": name, "note": "", "tags": []})
    assert response.status_code == 201, response.text
    return response.json()["database"]["key"]


def import_rows(client: TestClient, key: str, rows: list[tuple], file_name: str = "价目.xlsx") -> dict:
    response = client.post(
        f"/api/databases/{key}/import",
        files={"file": (file_name, price_workbook(rows),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def list_rows(client: TestClient, key: str, **params) -> dict:
    response = client.get("/api/history/rows", params={"database_key": key, **params})
    assert response.status_code == 200, response.text
    return response.json()


def bulk_edit(client: TestClient, key: str, **payload) -> dict:
    response = client.post("/api/history/bulk-edit", json={"database_key": key, **payload})
    assert response.status_code == 200, response.text
    return response.json()


def load_record(key: str, record_id: str) -> HistoryQuote:
    """直接从库文件读一行，用来核对派生列/版本号这类接口不返回的字段。"""
    session = dbmod.session_for_key(key)
    try:
        return session.get(HistoryQuote, record_id)
    finally:
        session.close()


# ---- 分页浏览 ---------------------------------------------------------------

def test_rows_endpoint_paginates_and_reports_total():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "分页浏览价目库")
        rows = [(f"分页测试器材{index:02d}", f"{index}00ml", "个", "分页品牌", "分页制造商", 10.0 + index)
                for index in range(1, 26)]
        assert import_rows(client, key, rows)["inserted"] == 25

        first = list_rows(client, key, page=1, page_size=10)
        assert first["total"] == 25
        assert first["pages"] == 3
        assert len(first["rows"]) == 10

        third = list_rows(client, key, page=3, page_size=10)
        assert len(third["rows"]) == 5

        # 分页不能重叠也不能漏：三页 id 合起来正好是全部 25 行
        ids = []
        for page in (1, 2, 3):
            ids.extend(item["id"] for item in list_rows(client, key, page=page, page_size=10)["rows"])
        assert len(set(ids)) == 25


def test_rows_endpoint_exposes_field_metadata_and_revision():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "字段元数据价目库")
        import_rows(client, key, [("字段元数据器材", "规格A", "台", "品牌A", "制造商A", 88.0)])

        body = list_rows(client, key)
        assert "name" in body["editable_fields"] and "price" in body["editable_fields"]
        assert "id" in body["readonly_fields"] and "revision" in body["readonly_fields"]
        assert body["rows"][0]["revision"] == 1
        # 派生列不能出现在可编辑列表里，否则前端会让人改到不该改的东西
        assert not set(body["editable_fields"]) & {"normalized_name", "data_quality"}


def test_rows_endpoint_sorts_by_price_desc():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "排序价目库")
        import_rows(client, key, [
            ("排序器材甲", "规格1", "个", "品牌", "制造商", 30.0),
            ("排序器材乙", "规格2", "个", "品牌", "制造商", 10.0),
            ("排序器材丙", "规格3", "个", "品牌", "制造商", 20.0),
        ])
        body = list_rows(client, key, sort="price", order="desc")
        assert [item["price"] for item in body["rows"]] == [30.0, 20.0, 10.0]


def test_rows_endpoint_filters_by_keyword():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "关键词价目库")
        import_rows(client, key, [
            ("关键词过滤量筒", "100ml", "个", "品牌", "制造商", 12.0),
            ("关键词过滤烧杯", "250ml", "个", "品牌", "制造商", 15.0),
        ])
        body = list_rows(client, key, q="量筒")
        assert body["total"] == 1
        assert body["rows"][0]["name"] == "关键词过滤量筒"


def test_rows_endpoint_returns_404_for_missing_database():
    with TestClient(app) as client:
        login(client)
        response = client.get("/api/history/rows", params={"database_key": "根本不存在"})
        assert response.status_code == 404


# ---- 修改 ------------------------------------------------------------------

def test_bulk_edit_update_recomputes_derived_columns_and_bumps_revision():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "派生列价目库")
        import_rows(client, key, [("派生列器材", "旧规格", "个", "品牌", "制造商", 50.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        result = bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 1,
            "changes": {"spec": "新规格 200ml", "price": 66.5, "manufacturer": "新制造商"},
        }])
        assert result["applied"]["updated"] == 1
        assert result["results"][0]["status"] == "ok"
        assert set(result["results"][0]["changed"]) == {"spec", "price", "manufacturer"}

        stored = load_record(key, record_id)
        assert stored.price == 66.5
        assert stored.manufacturer == "新制造商"
        # 派生列必须跟着业务列重算，否则治理页面的重复/冲突分组会用到旧口径
        assert stored.normalized_spec == "新规格200ml"
        assert stored.normalized_unit == "个"
        assert stored.revision == 2
        # 六个计分字段里 spec/model/unit/manufacturer/product_code/source_file
        # 有 spec+unit+manufacturer+source_file 四项非空
        assert abs(stored.data_quality - 4 / 6) < 1e-9


def test_bulk_edit_update_reports_unchanged_without_bumping_revision():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "空改价目库")
        import_rows(client, key, [("空改器材", "规格", "个", "品牌", "制造商", 10.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        result = bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 1, "changes": {"name": "空改器材"},
        }])
        assert result["applied"]["updated"] == 0
        assert result["unchanged"] == 1
        assert result["results"][0]["changed"] == []
        assert load_record(key, record_id).revision == 1


# ---- 版本冲突（乐观锁） -----------------------------------------------------

def test_bulk_edit_rejects_stale_revision_and_accepts_current_one():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "版本冲突价目库")
        import_rows(client, key, [("版本冲突器材", "规格", "个", "品牌", "制造商", 10.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        # 别人先改了一版
        assert bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 1, "changes": {"price": 20.0},
        }])["applied"]["updated"] == 1

        # 我拿着旧版本 1 去改 → 冲突，且回传当前版本号
        stale = bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 1, "changes": {"price": 999.0},
        }])
        assert stale["applied"]["updated"] == 0
        item = stale["results"][0]
        assert item["status"] == "conflict"
        assert item["reason"] == "revision_mismatch"
        assert item["conflict"]["actual_revision"] == 2
        assert item["conflict"]["expected_revision"] == 1
        assert load_record(key, record_id).price == 20.0, "冲突行不能被写入"

        # 用当前版本重试 → 成功
        retried = bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 2, "changes": {"price": 999.0},
        }])
        assert retried["applied"]["updated"] == 1
        assert load_record(key, record_id).price == 999.0


def test_bulk_edit_reports_missing_row_as_rejected():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "行缺失价目库")
        result = bulk_edit(client, key, updated=[{
            "id": "不存在的行id", "revision": 1, "changes": {"price": 1.0},
        }])
        assert result["rejected"] == 1
        assert result["results"][0]["reason"] == "missing"


# ---- 去重冲突：拒绝并给出位置 -----------------------------------------------

def test_bulk_edit_rejects_duplicate_create_and_returns_location():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "去重冲突价目库")
        import_rows(client, key, [("去重冲突器材", "规格X", "个", "品牌", "制造商", 45.0)],
                    file_name="去重来源本.xlsx")
        existing = list_rows(client, key)["rows"][0]

        result = bulk_edit(client, key, created=[{
            "changes": {"name": "去重冲突器材", "spec": "规格X", "unit": "个",
                        "manufacturer": "制造商", "price": 45.0},
        }])
        assert result["applied"]["created"] == 0
        assert result["skipped_duplicates"] == 1
        item = result["results"][0]
        assert item["status"] == "conflict"
        assert item["reason"] == "duplicate_key"
        # 位置要能定位到原行，前端据此跳转
        assert item["conflict"]["id"] == existing["id"]
        assert item["conflict"]["name"] == "去重冲突器材"
        assert "去重来源本.xlsx" in item["conflict"]["location"]
        assert "第2行" in item["conflict"]["location"]
        assert list_rows(client, key)["total"] == 1, "重复行不能被写入"


def test_bulk_edit_rejects_update_that_would_collide_with_another_row():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "改后撞车价目库")
        import_rows(client, key, [
            ("改后撞车甲", "规格1", "个", "品牌", "制造商", 10.0),
            ("改后撞车乙", "规格2", "个", "品牌", "制造商", 20.0),
        ])
        rows = list_rows(client, key)["rows"]
        first = next(item for item in rows if item["name"] == "改后撞车甲")
        second = next(item for item in rows if item["name"] == "改后撞车乙")

        # 去重键是（归一化名称, 参数, 单位, 制造商, 价格）五元组，必须整键撞上
        # 才算重复——只改价格不构成重复。
        result = bulk_edit(client, key, updated=[{
            "id": second["id"], "revision": 1,
            "changes": {"name": "改后撞车甲", "spec": "规格1", "price": 10.0},
        }])
        assert result["applied"]["updated"] == 0
        item = result["results"][0]
        assert item["reason"] == "duplicate_key"
        assert item["conflict"]["id"] == first["id"]
        stored = load_record(key, second["id"])
        assert stored.name == "改后撞车乙" and stored.spec == "规格2"


def test_bulk_edit_allows_updating_a_row_in_place():
    """同一行只改价格不算撞自己——去重键变化后不能和自己判重。"""
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "自身不撞价目库")
        import_rows(client, key, [("自身不撞器材", "规格", "个", "品牌", "制造商", 10.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        result = bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 1, "changes": {"price": 12.0},
        }])
        assert result["applied"]["updated"] == 1
        assert result["results"][0]["status"] == "ok"


# ---- 新增 / 删除 ------------------------------------------------------------

def test_bulk_edit_creates_row_with_manual_id_and_traced_source():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "新增行价目库")
        result = bulk_edit(client, key, created=[{
            "changes": {"name": "新增行器材", "spec": "300ml", "unit": "个",
                        "manufacturer": "新增制造商", "price": 33.0},
        }])
        assert result["applied"]["created"] == 1
        new_id = result["results"][0]["id"]
        assert new_id.startswith("manual:")

        body = list_rows(client, key)
        assert body["total"] == 1
        row = body["rows"][0]
        assert row["name"] == "新增行器材"
        # 新增行必须有可追溯来源，否则审计里查不到它是从哪来的
        assert row["source_file"] == "表格编辑"
        assert row["revision"] == 1


def test_bulk_edit_deletes_row_and_writes_audit_event():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "删除行价目库")
        import_rows(client, key, [("删除行器材", "规格", "个", "品牌", "制造商", 10.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        result = bulk_edit(client, key, deleted=[{"id": record_id, "revision": 1}])
        assert result["applied"]["deleted"] == 1
        assert list_rows(client, key)["total"] == 0
        assert load_record(key, record_id) is None

        session = dbmod.session_for_key(key)
        try:
            actions = set(session.scalars(select(AuditEvent.action)).all())
        finally:
            session.close()
        assert "history.bulk_edit" in actions


def test_bulk_edit_delete_also_frees_the_dedup_key():
    """删掉旧行后，同一条目应当可以再新增回来（删除要摘掉去重键）。"""
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "删后可加价目库")
        import_rows(client, key, [("删后可加器材", "规格", "个", "品牌", "制造商", 10.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]
        assert bulk_edit(client, key, deleted=[{"id": record_id, "revision": 1}])["applied"]["deleted"] == 1

        again = bulk_edit(client, key, created=[{
            "changes": {"name": "删后可加器材", "spec": "规格", "unit": "个",
                        "manufacturer": "制造商", "price": 10.0},
        }])
        assert again["applied"]["created"] == 1
        assert again["skipped_duplicates"] == 0


# ---- 校验：逐行拒绝，不影响其它行 -------------------------------------------

def test_bulk_edit_rejects_invalid_rows_but_applies_valid_ones():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "逐行校验价目库")
        result = bulk_edit(client, key, created=[
            {"changes": {"name": "有效器材", "spec": "规格", "unit": "个", "price": 10.0}},
            {"changes": {"name": "   ", "spec": "规格", "unit": "个", "price": 10.0}},
            {"changes": {"name": "负价器材", "spec": "规格", "unit": "个", "price": -5.0}},
            {"changes": {"name": "非数器材", "spec": "规格", "unit": "个", "price": "abc"}},
            {"changes": {"name": "缺价器材", "spec": "规格", "unit": "个", "price": None}},
        ])
        assert result["applied"]["created"] == 1
        assert result["rejected"] == 4
        reasons = [item["reason"] for item in result["results"][1:]]
        assert reasons == ["invalid"] * 4
        assert list_rows(client, key)["total"] == 1, "只有合法行落库"


def test_bulk_edit_rejects_readonly_field_changes():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "只读字段价目库")
        import_rows(client, key, [("只读字段器材", "规格", "个", "品牌", "制造商", 10.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        result = bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 1,
            "changes": {"source_file": "偷偷改来源", "price": 11.0},
        }])
        assert result["rejected"] == 1
        assert "source_file" in result["results"][0]["message"]
        stored = load_record(key, record_id)
        assert stored.source_file != "偷偷改来源"
        assert stored.price == 10.0


def test_bulk_edit_enforces_single_batch_row_limit():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "批量上限价目库")
        too_many = [{"changes": {"name": f"批量上限{index}", "price": 1.0}} for index in range(2001)]
        response = client.post("/api/history/bulk-edit", json={"database_key": key, "created": too_many})
        assert response.status_code == 400
        assert "2000" in response.json()["detail"]


def test_bulk_edit_requires_admin():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "权限校验价目库")
        import_rows(client, key, [("权限校验器材", "规格", "个", "品牌", "制造商", 10.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        client.post("/api/auth/logout")
        login(client, "quote", "quote123")
        response = client.post("/api/history/bulk-edit", json={
            "database_key": key,
            "updated": [{"id": record_id, "revision": 1, "changes": {"price": 1.0}}],
        })
        assert response.status_code == 403
        assert load_record(key, record_id).price == 10.0


# ---- 备份 ------------------------------------------------------------------

def test_bulk_edit_creates_backup_before_writing():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "备份价目库")
        import_rows(client, key, [("备份器材", "规格", "个", "品牌", "制造商", 10.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        result = bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 1, "changes": {"price": 77.0},
        }])
        assert result["backup"].startswith(f"{key}_")
        backup = dbmod.DATA_DIR / "_backup" / result["backup"]
        assert backup.exists() and backup.stat().st_size > 0

        # 备份里应当是改动**之前**的内容，否则回退没有意义
        import sqlite3

        connection = sqlite3.connect(str(backup))
        try:
            price = connection.execute(
                "SELECT price FROM history_quotes WHERE id = ?", (record_id,)
            ).fetchone()[0]
        finally:
            connection.close()
        assert price == 10.0


# ---- 端到端：改完价目，报价必须用新价 ---------------------------------------

def test_edited_price_takes_effect_in_new_quote():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "生效验证价目库")
        product = "生效验证专用烧杯"
        import_rows(client, key, [(product, "250ml 玻璃", "个", "生效品牌", "生效制造商", 100.0)])
        record_id = list_rows(client, key)["rows"][0]["id"]

        def run_job() -> dict:
            created = client.post(
                "/api/quote-jobs",
                data={"requested_option_count": "1", "database_key": key},
                files={"file": ("生效询价.xlsx", inquiry_workbook(product, "250ml 玻璃", "个"),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert created.status_code == 202, created.text
            job = created.json()
            lines = client.get(f"/api/quote-jobs/{job['id']}/lines").json()
            return client.get(f"/api/quote-lines/{lines['items'][0]['id']}").json()

        before = run_job()
        assert before["options"], "导入后应当匹配出候选"
        assert before["options"][0]["record"]["price"] == 100.0

        saved = bulk_edit(client, key, updated=[{
            "id": record_id, "revision": 1,
            "changes": {"price": 258.0, "manufacturer": "生效新制造商"},
        }])
        assert saved["applied"]["updated"] == 1

        after = run_job()
        assert after["options"], "改价后仍应匹配出候选"
        primary = sorted(after["options"], key=lambda item: item["rank"])[0]
        assert primary["record"]["price"] == 258.0
        assert primary["record"]["manufacturer"] == "生效新制造商"


# ---- 老库迁移 ---------------------------------------------------------------

def test_legacy_database_gets_revision_column_before_editing(tmp_path):
    """老价目库没有 revision 列；不补迁移的话打开/保存会直接报 no such column。"""
    import sqlite3

    legacy = tmp_path / "legacy_price.db"
    connection = sqlite3.connect(str(legacy))
    try:
        connection.execute(
            "CREATE TABLE history_quotes ("
            "id VARCHAR(180) PRIMARY KEY, source_file VARCHAR(300), source_sheet VARCHAR(200), "
            "source_row INTEGER, name VARCHAR(500), normalized_name VARCHAR(500), spec TEXT, "
            "normalized_spec TEXT, product_code VARCHAR(200), normalized_product_code VARCHAR(200), "
            "model VARCHAR(300), normalized_model VARCHAR(300), brand VARCHAR(300), "
            "manufacturer VARCHAR(500), unit VARCHAR(80), normalized_unit VARCHAR(80), "
            "quantity FLOAT, price FLOAT, quote_date VARCHAR(40), source_priority INTEGER, "
            "data_quality FLOAT, created_at DATETIME)"
        )
        connection.execute(
            "INSERT INTO history_quotes (id, name, normalized_name, price) VALUES (?, ?, ?, ?)",
            ("legacy:1", "老库器材", "老库器材", 12.0),
        )
        connection.commit()
    finally:
        connection.close()

    dbmod.ensure_schema_for_path(legacy)

    connection = sqlite3.connect(str(legacy))
    try:
        columns = {row[1] for row in connection.execute("PRAGMA table_info(history_quotes)")}
        revision = connection.execute(
            "SELECT revision FROM history_quotes WHERE id = 'legacy:1'"
        ).fetchone()[0]
    finally:
        connection.close()
    assert "revision" in columns
    assert revision == 1


# ---- 打开本地 Excel：解析比对（不写库） -------------------------------------

HEADER_LABELS = {
    "name": "产品名称", "spec": "参数", "model": "型号", "brand": "品牌",
    "manufacturer": "制造商", "unit": "单位", "product_code": "产品编码", "price": "单价",
}


def build_workbook(columns: list[str], rows: list[tuple], title: str = "价目") -> bytes:
    """按指定列（英文字段名顺序）构建价目本；rows 的元组按同一顺序给值。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append([HEADER_LABELS[column] for column in columns])
    for row in rows:
        sheet.append(list(row))
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_workbook_sheets(sheets: list[tuple[str, list[str], list[tuple]]]) -> bytes:
    """多 sheet 价目本：每项为 (sheet 名, 列, 行)。"""
    workbook = Workbook()
    for index, (title, columns, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append([HEADER_LABELS[column] for column in columns])
        for row in rows:
            sheet.append(list(row))
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def post_workbook(client: TestClient, key: str, body: bytes, file_name: str = "打开价目.xlsx"):
    return client.post(
        "/api/history/open-file",
        data={"database_key": key},
        files={"file": (file_name, body,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def open_file(client: TestClient, key: str, columns: list[str], rows: list[tuple],
              file_name: str = "打开价目.xlsx") -> dict:
    response = post_workbook(client, key, build_workbook(columns, rows), file_name)
    assert response.status_code == 200, response.text
    return response.json()


def by_seq(payload: dict) -> dict[int, dict]:
    return {item["seq"]: item for item in payload["rows"]}


ALL_COLUMNS = ["name", "spec", "model", "brand", "manufacturer", "unit", "product_code", "price"]


def test_identity_key_is_the_prefix_of_the_dedup_key():
    """身份键必须正好是去重键的前四项——两个函数分居两个模块，靠这个断言锁住关系。"""
    from app.services import history_dedup_key

    for name, spec, unit, maker, price in [
        ("打孔器", "四支空芯钻头", "套", "某某厂", 6.6),
        ("烧杯", "250ml 玻璃", "个", "", 12.0),
        ("量筒", "１００ｍｌ", " 个 ", "Ａ厂", 0),
    ]:
        assert history_dedup_key(name, spec, unit, maker, price)[:4] == \
            history_rows.identity_key_values(name, spec, unit, maker)


def test_open_file_classifies_create_update_and_unchanged():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "打开文件分类价目库")
        import_rows(client, key, [
            ("打开文件甲", "规格A", "个", "原品牌", "原制造商", 10.0),
            ("打开文件乙", "规格B", "个", "原品牌", "原制造商", 20.0),
        ])

        payload = open_file(client, key, ALL_COLUMNS, [
            # 第 1 行要"完全一致"：import_rows 用的价目本没有型号/产品编码列，
            # 库里这两列是空的，文件若填了值就构成一次真实改动（见第 2 行）。
            ("打开文件甲", "规格A", "", "原品牌", "原制造商", "个", "", 10.0),
            ("打开文件乙", "规格B", "M-2", "新品牌", "原制造商", "个", "C-2", 25.0),   # 改价+改品牌+补型号
            ("打开文件丙", "规格C", "M-3", "新品牌", "新制造商", "个", "C-3", 30.0),   # 库里没有
        ])

        assert payload["summary"] == {
            "create": 1, "update": 1, "unchanged": 1,
            "ambiguous": 0, "duplicate": 0, "invalid": 0,
        }
        rows = by_seq(payload)
        assert rows[1]["action"] == "unchanged"
        assert rows[1]["changed"] == []

        update = rows[2]
        assert update["action"] == "update"
        # 匹配用到的身份字段要回传给前端，好把"凭什么说这两条是同一条"讲清楚
        assert update["match_fields"] == ["name", "spec", "unit", "manufacturer"]
        assert update["match_labels"] == ["名称", "参数", "单位", "制造商"]
        assert set(update["changed"]) == {"brand", "model", "price", "product_code"}
        # 合并后的值 = 库里现值 + 文件提供的字段；没被文件覆盖的字段保留原值
        assert update["values"]["price"] == 25.0
        assert update["values"]["manufacturer"] == "原制造商"
        assert update["current"]["price"] == 20.0
        assert update["current"]["brand"] == "原品牌"
        # 目标带 revision，前端据此做乐观锁
        assert update["target"]["revision"] == 1
        # 定位指向"库里这条是从哪个文件哪一行来的"：导入价目本第 1 行是表头，
        # 所以第二件器材落在源文件第 3 行
        assert "价目.xlsx" in update["target"]["location"]
        assert "第3行" in update["target"]["location"]

        assert rows[3]["action"] == "create"
        assert rows[3]["values"]["price"] == 30.0


def test_open_file_price_only_change_is_an_update_not_a_create():
    """身份键不含价格：只改价必须认成"修改这一行"，不能新增出一条重复记录。"""
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "改价识别价目库")
        import_rows(client, key, [("只改价器材", "规格A", "个", "品牌", "制造商", 10.0)])

        payload = open_file(client, key, ALL_COLUMNS, [
            ("只改价器材", "规格A", "", "品牌", "制造商", "个", "", 18.0),
        ])
        item = payload["rows"][0]
        assert item["action"] == "update"
        assert item["changed"] == ["price"]
        assert item["values"]["price"] == 18.0
        assert payload["summary"]["create"] == 0
        # 预览阶段不落库
        assert list_rows(client, key)["total"] == 1


def test_open_file_never_writes_to_the_database():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "只读预览价目库")
        import_rows(client, key, [("只读预览器材", "规格A", "个", "品牌", "制造商", 10.0)])
        before = list_rows(client, key)

        open_file(client, key, ALL_COLUMNS, [
            ("只读预览器材", "规格A", "", "品牌", "制造商", "个", "", 99.0),
            ("只读预览新增", "规格B", "", "品牌", "制造商", "个", "", 5.0),
        ])

        after = list_rows(client, key)
        assert after["total"] == before["total"] == 1
        assert after["rows"][0]["price"] == 10.0
        assert after["rows"][0]["revision"] == 1


def test_open_file_without_brand_column_preserves_existing_brand():
    """文件没有品牌列 = 不关心品牌，更新时必须保留库里的原值。"""
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "缺列保留价目库")
        import_rows(client, key, [("缺列保留器材", "规格A", "个", "原品牌", "原制造商", 10.0)])

        columns = ["name", "spec", "unit", "price"]
        payload = open_file(client, key, columns, [("缺列保留器材", "规格A", "个", 15.0)])
        item = payload["rows"][0]
        assert item["action"] == "update"
        assert item["changed"] == ["price"]
        assert item["values"]["brand"] == "原品牌"
        assert item["values"]["manufacturer"] == "原制造商"


def test_open_file_with_empty_brand_cell_clears_the_field():
    """文件有品牌列但这一行空着 = 用户想清空，必须真的清空。"""
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "空列清空价目库")
        import_rows(client, key, [("空列清空器材", "规格A", "个", "原品牌", "原制造商", 10.0)])

        payload = open_file(client, key, ALL_COLUMNS, [
            ("空列清空器材", "规格A", "", "", "原制造商", "个", "", 10.0),
        ])
        item = payload["rows"][0]
        assert item["action"] == "update"
        assert item["changed"] == ["brand"]
        assert item["values"]["brand"] == ""


def test_open_file_flags_ambiguous_and_offers_candidates():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "多条同身份价目库")
        # 去重键含价格，所以同身份不同价可以共存
        import_rows(client, key, [
            ("同身份器材", "规格A", "个", "品牌", "制造商", 10.0),
            ("同身份器材", "规格A", "个", "品牌", "制造商", 20.0),
        ])

        payload = open_file(client, key, ALL_COLUMNS, [
            ("同身份器材", "规格A", "", "品牌", "制造商", "个", "", 30.0),
        ])
        item = payload["rows"][0]
        assert item["action"] == "ambiguous"
        assert len(item["targets"]) == 2
        # 这行自己也要带文件里的值：前端靠它显示"是哪件器材要指定"，只有行号没法认。
        assert item["values"]["name"] == "同身份器材"
        assert item["values"]["price"] == 30.0
        # 每个候选都预先算好了合并结果，前端选了就能直接用
        for target in item["targets"]:
            assert target["changed"] == ["price"]
            assert target["values"]["price"] == 30.0
            assert target["current"]["price"] in (10.0, 20.0)
            assert target["revision"] == 1
        assert payload["summary"]["ambiguous"] == 1


def test_open_file_flags_duplicate_rows_within_the_file():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "文件内重复价目库")
        payload = open_file(client, key, ALL_COLUMNS, [
            ("文件内重复器材", "规格A", "", "品牌", "制造商", "个", "", 10.0),
            ("文件内重复器材", "规格A", "", "品牌", "制造商", "个", "", 10.0),
        ])
        rows = by_seq(payload)
        assert rows[1]["action"] == "create"
        assert rows[2]["action"] == "duplicate"
        assert "第 1 行" in rows[2]["reason"]
        assert payload["summary"] == {
            "create": 1, "update": 0, "unchanged": 0,
            "ambiguous": 0, "duplicate": 1, "invalid": 0,
        }


def test_open_file_without_manufacturer_column_matches_on_the_rest():
    """文件没有制造商列 = 算不出完整身份，退化成名称+参数+单位，仍要认成更新。

    这是真实踩过的坑：身份键含制造商，文件缺这一列时只能算成空串，拿空串去和
    库里"制造商=某厂"的行比永远比不上，"更新"被误判成"新增"，一保存就多出一条
    重复记录。
    """
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "缺制造商列价目库")
        import_rows(client, key, [("缺制造商器材", "规格A", "个", "品牌", "某厂", 10.0)])

        payload = open_file(client, key, ["name", "spec", "unit", "price"],
                            [("缺制造商器材", "规格A", "个", 18.0)])
        item = payload["rows"][0]
        assert item["action"] == "update"
        assert item["match_fields"] == ["name", "spec", "unit"]
        assert item["changed"] == ["price"]
        assert item["values"]["manufacturer"] == "某厂"  # 没提供 → 保留库里的原值
        assert payload["summary"]["create"] == 0


def test_open_file_loose_match_reports_ambiguous_when_several_candidates():
    """没有制造商列时匹配更松：同名同规格不同制造商的两条都会命中 → 交给人工指定。"""
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "松匹配歧义价目库")
        import_rows(client, key, [
            ("松匹配器材", "规格A", "个", "品牌", "甲厂", 10.0),
            ("松匹配器材", "规格A", "个", "品牌", "乙厂", 20.0),
        ])

        payload = open_file(client, key, ["name", "spec", "unit", "price"],
                            [("松匹配器材", "规格A", "个", 30.0)])
        item = payload["rows"][0]
        assert item["action"] == "ambiguous"
        assert item["match_fields"] == ["name", "spec", "unit"]
        assert len(item["targets"]) == 2
        assert {target["current"]["manufacturer"] for target in item["targets"]} == {"甲厂", "乙厂"}


def test_open_file_flags_two_sheets_pointing_at_the_same_record():
    """两个 sheet 列不同 → 身份键不同，但可能指向库里同一条；第二条必须判重复。

    否则保存时两行会去改同一条记录，第二行必然撞乐观锁，用户看到的是莫名其妙的
    "版本冲突"而不是"你在文件里写了两次同一件东西"。
    """
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "跨sheet同目标价目库")
        import_rows(client, key, [("跨sheet器材", "规格A", "个", "品牌", "甲厂", 10.0)])

        body = build_workbook_sheets([
            ("无制造商列", ["name", "spec", "unit", "price"], [("跨sheet器材", "规格A", "个", 20.0)]),
            ("带制造商列", ["name", "spec", "manufacturer", "unit", "price"],
             [("跨sheet器材", "规格A", "甲厂", "个", 30.0)]),
        ])
        response = post_workbook(client, key, body)
        assert response.status_code == 200, response.text
        payload = response.json()
        rows = by_seq(payload)
        assert rows[1]["action"] == "update"
        assert rows[2]["action"] == "duplicate"
        assert "同一条记录" in rows[2]["reason"]
        assert payload["summary"]["duplicate"] == 1


def test_open_file_marks_rows_without_a_valid_price_as_invalid():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "无价行价目库")
        payload = open_file(client, key, ALL_COLUMNS, [
            ("无价行器材", "规格A", "", "品牌", "制造商", "个", "", None),
            ("零价行器材", "规格B", "", "品牌", "制造商", "个", "", 0),
        ])
        assert payload["summary"]["invalid"] == 2
        for item in payload["rows"]:
            assert item["action"] == "invalid"
            assert "单价" in item["reason"]


def test_open_file_rejects_workbook_without_any_price_column():
    """整本都没有价格列 → 是文件级错误（400），不是逐行 invalid。

    一行价格都没有的价目本没有任何可用信息，报错让人换文件比列 5000 行 invalid
    更有用。
    """
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "整本无价列价目库")
        response = post_workbook(client, key, build_workbook(["name", "spec", "unit"],
                                                             [("无价列器材", "规格A", "个")]))
        assert response.status_code == 400
        assert "未找到价格列" in response.json()["detail"]


def test_open_file_marks_sheet_without_price_column_as_invalid():
    """多 sheet 里只有个别 sheet 没价格列 → 该 sheet 逐行 invalid，不拖累整本。"""
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "无价列sheet价目库")
        body = build_workbook_sheets([
            ("有价列", ["name", "spec", "unit", "price"], [("有价列器材", "规格A", "个", 5.0)]),
            ("无价列", ["name", "spec", "unit"], [("无价列器材", "规格A", "个")]),
        ])
        response = post_workbook(client, key, body)
        assert response.status_code == 200, response.text
        payload = response.json()
        rows = by_seq(payload)
        assert rows[1]["action"] == "create"
        assert rows[2]["action"] == "invalid"
        assert "没有价格列" in rows[2]["reason"]
        assert payload["summary"] == {
            "create": 1, "update": 0, "unchanged": 0,
            "ambiguous": 0, "duplicate": 0, "invalid": 1,
        }


def test_open_file_enforces_row_limit(monkeypatch):
    from app import main as main_module

    with TestClient(app) as client:
        login(client)
        key = create_database(client, "行数上限价目库")
        monkeypatch.setattr(main_module, "OPEN_FILE_MAX_ROWS", 1)
        response = client.post(
            "/api/history/open-file",
            data={"database_key": key},
            files={"file": ("超限.xlsx", build_workbook(ALL_COLUMNS, [
                ("超限甲", "规格A", "", "品牌", "制造商", "个", "", 1.0),
                ("超限乙", "规格B", "", "品牌", "制造商", "个", "", 2.0),
            ]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 400
        assert "上限" in response.json()["detail"]


def test_open_file_requires_admin_and_existing_database():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "打开文件权限价目库")
        body = build_workbook(ALL_COLUMNS, [("权限器材", "规格A", "", "品牌", "制造商", "个", "", 1.0)])

        missing = client.post(
            "/api/history/open-file",
            data={"database_key": "不存在的库"},
            files={"file": ("x.xlsx", body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert missing.status_code == 404

        client.post("/api/auth/logout")
        login(client, "quote", "quote123")
        forbidden = client.post(
            "/api/history/open-file",
            data={"database_key": key},
            files={"file": ("x.xlsx", body, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert forbidden.status_code == 403


def test_open_file_rejects_non_excel():
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "非Excel价目库")
        response = client.post(
            "/api/history/open-file",
            data={"database_key": key},
            files={"file": ("备注.txt", b"not excel", "text/plain")},
        )
        assert response.status_code == 400
        assert ".xlsx" in response.json()["detail"]


def test_open_file_preview_then_save_round_trip():
    """完整闭环：打开文件 → 应用 → 保存 → 库里真的是文件里的值。"""
    with TestClient(app) as client:
        login(client)
        key = create_database(client, "打开保存闭环价目库")
        import_rows(client, key, [
            ("闭环器材甲", "规格A", "个", "品牌", "制造商", 10.0),
            ("闭环器材乙", "规格B", "个", "品牌", "制造商", 20.0),
        ])

        payload = open_file(client, key, ALL_COLUMNS, [
            ("闭环器材甲", "规格A", "", "品牌", "制造商", "个", "", 12.5),
            ("闭环器材丙", "规格C", "", "品牌", "制造商", "个", "", 33.0),
        ])
        rows = by_seq(payload)
        assert rows[1]["action"] == "update"
        assert rows[2]["action"] == "create"

        saved = bulk_edit(
            client, key,
            updated=[{
                "id": rows[1]["target"]["id"],
                "revision": rows[1]["target"]["revision"],
                "changes": rows[1]["values"],
            }],
            created=[{"changes": rows[2]["values"]}],
        )
        assert saved["applied"] == {"created": 1, "updated": 1, "deleted": 0}
        assert saved["rejected"] == 0

        final = {row["name"]: row for row in list_rows(client, key, page_size=100)["rows"]}
        assert final["闭环器材甲"]["price"] == 12.5
        assert final["闭环器材甲"]["revision"] == 2
        assert final["闭环器材丙"]["price"] == 33.0
        assert final["闭环器材丙"]["source_file"] == "表格编辑"
