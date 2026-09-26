"""数据资产中心：库注册表、CRUD、搜索/最近常用，以及任务级绑定（B1）。"""

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app import registry
from app.database import system_db_key
from app.registry import normalize_name


def login(client: TestClient, username: str = "admin", password: str = "admin123"):
    response = client.post("/api/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text


def price_workbook(name: str, spec: str, unit: str, price: float, brand: str, manufacturer: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "价目"
    sheet.append(["序号", "产品名称", "参数", "单位", "品牌", "制造商", "含税单价"])
    sheet.append([1, name, spec, unit, brand, manufacturer, price])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def inquiry_workbook(name: str, spec: str, unit: str, quantity: int) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "询价单"
    sheet.append(["序号", "产品名称", "参数", "型号", "单位", "数量"])
    sheet.append([1, name, spec, "", unit, quantity])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def create_database(client: TestClient, name: str, note: str = "", tags: list[str] | None = None) -> dict:
    response = client.post("/api/databases", json={"name": name, "note": note, "tags": tags or []})
    assert response.status_code == 201, response.text
    return response.json()["database"]


def database_keys(client: TestClient) -> dict[str, dict]:
    return {item["name"]: item for item in client.get("/api/databases").json()["databases"]}


# ---- 中文名与 CRUD ---------------------------------------------------------

def test_database_create_uses_chinese_display_name_and_ascii_key():
    with TestClient(app) as client:
        login(client)
        entry = create_database(client, "赛特尔25年", note="高中理化生主价目本", tags=["普教", "高中"])
        # 显示名保留中文，文件键是 ASCII（键名分离，重命名不动文件）
        assert entry["name"] == "赛特尔25年"
        assert entry["key"].isascii() and entry["key"].isidentifier() or entry["key"].replace("_", "").isalnum()
        assert entry["note"] == "高中理化生主价目本"
        assert entry["tags"] == ["普教", "高中"]
        assert entry["exists"] is True

        listed = database_keys(client)
        assert listed["赛特尔25年"]["key"] == entry["key"]


def test_database_rename_keeps_key_and_file():
    with TestClient(app) as client:
        login(client)
        entry = create_database(client, "待改名库")
        renamed = client.patch(f"/api/databases/{entry['key']}", json={"name": "改名后的中文库"})
        assert renamed.status_code == 200, renamed.text
        body = renamed.json()["database"]
        assert body["name"] == "改名后的中文库"
        # 重命名只改显示名，不动文件键
        assert body["key"] == entry["key"]
        assert "改名后的中文库" in database_keys(client)
        assert "待改名库" not in database_keys(client)


def test_database_name_uniqueness_ignores_spaces_and_width():
    with TestClient(app) as client:
        login(client)
        create_database(client, "普教清单")
        # 全角/半角、空格差异视为同一个名字
        conflict = client.post("/api/databases", json={"name": "普教 清单", "note": "", "tags": []})
        assert conflict.status_code == 400, conflict.text
        assert "已存在" in conflict.json()["detail"]
        assert normalize_name("普教 清单") == normalize_name("普教清单")


def test_database_rejects_illegal_name():
    with TestClient(app) as client:
        login(client)
        for bad in ("", "   ", "含/斜杠", "含:冒号", "x" * 41):
            response = client.post("/api/databases", json={"name": bad, "note": "", "tags": []})
            assert response.status_code in (400, 422), f"{bad!r} 应被拒绝: {response.text}"


def test_database_trash_restore_and_purge():
    with TestClient(app) as client:
        login(client)
        entry = create_database(client, "回收站测试库")
        key = entry["key"]

        deleted = client.delete(f"/api/databases/{key}")
        assert deleted.status_code == 200, deleted.text
        payload = client.get("/api/databases").json()
        assert key not in {item["key"] for item in payload["databases"]}
        assert key in {item["key"] for item in payload["trashed"]}

        restored = client.post(f"/api/databases/{key}/restore")
        assert restored.status_code == 200, restored.text
        assert key in {item["key"] for item in client.get("/api/databases").json()["databases"]}

        # 彻底删除必须回填正确的库名
        assert client.delete(f"/api/databases/{key}").status_code == 200
        wrong = client.post(f"/api/databases/{key}/purge", json={"confirm_name": "随便一个名字"})
        assert wrong.status_code == 400
        right = client.post(f"/api/databases/{key}/purge", json={"confirm_name": "回收站测试库"})
        assert right.status_code == 200, right.text
        assert key not in {item["key"] for item in client.get("/api/databases").json()["trashed"]}


def test_system_database_cannot_be_deleted():
    """系统库是"默认库"的唯一来源，不允许删掉它——否则未指定库的读取路径全废。"""
    with TestClient(app) as client:
        login(client)
        system = client.get("/api/databases").json()["system"]
        assert system, "测试环境应有系统库"
        response = client.delete(f"/api/databases/{system}")
        assert response.status_code == 400
        assert "不能删除" in response.json()["detail"]
        # 系统库仍然在列表里，且被标记出来
        listed = client.get("/api/databases").json()["databases"]
        entry = next(item for item in listed if item["key"] == system)
        assert entry["is_system"] is True
        assert sum(item["is_system"] for item in listed) == 1, "有且只有一个系统库"


def test_system_display_name_blocks_duplicate_create():
    """系统库占用的不只是文件键，还有它的显示名。

    回归点：系统库的显示名从"文件名键"改成人能认的中文名之后，这个名字就成了
    一个真实被占用的库名——再建同名库应当 400，而不是悄悄多出一个同名条目。
    """
    with TestClient(app) as client:
        login(client)
        listed = client.get("/api/databases").json()["databases"]
        system_name = next(item["name"] for item in listed if item["is_system"])
        conflict = client.post("/api/databases", json={"name": system_name, "note": "", "tags": []})
        assert conflict.status_code == 400, conflict.text
        assert "已存在" in conflict.json()["detail"]


def test_activation_endpoints_are_gone():
    """改造后不再有"激活/切换库"——旧接口必须彻底消失，避免前端误用。"""
    with TestClient(app) as client:
        login(client)
        entry = create_database(client, "切换接口测试库")
        assert client.post(f"/api/databases/{entry['key']}/activate").status_code == 404
        # /api/databases/switch 被 /api/databases/{db_key}（PATCH/DELETE）吃掉路径，
        # 所以是 405 而不是 404——两种都表示"没有这个 POST 接口"。
        assert client.post("/api/databases/switch", json={"name": entry["name"]}).status_code in (404, 405)


def test_database_write_requires_admin_but_read_is_open():
    with TestClient(app) as client:
        login(client)
        entry = create_database(client, "权限测试库")
        client.post("/api/auth/logout")

        login(client, "quote", "quote123")
        # 报价员只读：可以看列表和搜索
        assert client.get("/api/databases").status_code == 200
        assert client.get("/api/databases/search", params={"limit": 5}).status_code == 200
        # 写操作全部 403
        assert client.post("/api/databases", json={"name": "越权库", "note": "", "tags": []}).status_code == 403
        assert client.patch(f"/api/databases/{entry['key']}", json={"name": "越权改名"}).status_code == 403
        assert client.delete(f"/api/databases/{entry['key']}").status_code == 403
        assert client.post(f"/api/databases/{entry['key']}/import", files={"file": ("a.xlsx", b"x", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).status_code == 403


# ---- 搜索与最近常用 --------------------------------------------------------

def test_search_returns_recent_five_by_default():
    with TestClient(app) as client:
        login(client)
        names = [f"最近常用{i}号库" for i in range(1, 8)]
        keys = {name: create_database(client, name)["key"] for name in names}

        # 用"导入价目本"标记使用（touch_use），顺序即最近使用顺序
        for name in ("最近常用1号库", "最近常用5号库", "最近常用3号库"):
            imported = client.post(
                f"/api/databases/{keys[name]}/import",
                files={"file": (f"{name}.xlsx", price_workbook(name, "规格", "台", 10.0, "品牌", "厂商"),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert imported.status_code == 200, imported.text
            assert imported.json()["inserted"] == 1

        recent = client.get("/api/databases/search", params={"limit": 5}).json()["databases"]
        assert len(recent) == 5, "默认只展示最近常用的 5 个"
        assert recent[0]["name"] == "最近常用3号库", "最后使用的排在最前"
        used = [item["name"] for item in recent]
        assert used[:3] == ["最近常用3号库", "最近常用5号库", "最近常用1号库"]

        # 导入后记录数应当已更新
        assert recent[0]["history_count"] == 1


def test_recent_five_prefers_newly_created_over_alphabetical():
    """新建但尚未使用过的库也应按"新"排前，而不是按名称字母序。

    回归点：旧实现只用 ``last_used_at`` 排序，未使用过的库该列全是 NULL，一起掉到
    末尾后按名称排序——于是"最近常用 5 个"退化成"名称前 5 个"，用户刚建好的库
    反而不出现在选择列表里。这里特意按升序命名（1..7 号），使"按名称"与"按时间"
    的结果不同：按时间应为 7/6/5/4/3，按名称则为 1/2/3/4/5。
    """
    with TestClient(app) as client:
        login(client)
        prefix = "新鲜度"
        for i in range(1, 8):
            create_database(client, f"{prefix}{i}号库")

        recent = client.get("/api/databases/search", params={"limit": 5}).json()["databases"]
        names = [item["name"] for item in recent]
        assert len(names) == 5, "默认只展示最近常用的 5 个"
        assert names[0] == f"{prefix}7号库", f"最后创建的应排最前，实际 {names}"
        assert set(names) == {f"{prefix}{i}号库" for i in range(3, 8)}, f"应取最新的 5 个，实际 {names}"
        # 明确排除"按名称取前 5"的退化行为
        assert f"{prefix}1号库" not in names and f"{prefix}2号库" not in names


def test_search_matches_chinese_substring():
    with TestClient(app) as client:
        login(client)
        # 用唯一前缀，避免与同会话其它测试创建的库互相干扰
        create_database(client, "检索专用高中理化库")
        create_database(client, "检索专用小学科学库")
        create_database(client, "检索专用普教清单库")

        hit = client.get("/api/databases/search", params={"q": "检索专用高中"}).json()["databases"]
        assert [item["name"] for item in hit] == ["检索专用高中理化库"]

        fuzzy = client.get("/api/databases/search", params={"q": "检索专用普教"}).json()["databases"]
        assert [item["name"] for item in fuzzy] == ["检索专用普教清单库"]

        none = client.get("/api/databases/search", params={"q": "绝无仅有的关键词xyz"}).json()["databases"]
        assert none == []


# ---- 任务级绑定（B1） ------------------------------------------------------

def test_quote_job_binds_selected_database_end_to_end():
    with TestClient(app) as client:
        login(client)
        entry = create_database(client, "绑定测试价目库")
        key = entry["key"]

        product = "绑定测试专用量筒"
        imported = client.post(
            f"/api/databases/{key}/import",
            files={"file": ("绑定测试.xlsx",
                            price_workbook(product, "100ml 玻璃", "个", 23.5, "绑定品牌", "绑定制造商"),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert imported.status_code == 200, imported.text
        assert imported.json()["inserted"] == 1

        created = client.post(
            "/api/quote-jobs",
            data={"requested_option_count": "3", "database_key": key},
            files={"file": ("绑定询价.xlsx", inquiry_workbook(product, "100ml 玻璃", "个", 2),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert created.status_code == 202, created.text
        job = created.json()
        # 任务记录了绑定库与当时的显示名快照
        assert job["database_key"] == key
        assert job["database_name"] == "绑定测试价目库"

        detail = client.get(f"/api/quote-jobs/{job['id']}").json()
        assert detail["database_key"] == key

        lines = client.get(f"/api/quote-jobs/{job['id']}/lines").json()
        line = client.get(f"/api/quote-lines/{lines['items'][0]['id']}").json()
        assert line["options"], "绑定库中的价目应当匹配出候选"
        primary = sorted(line["options"], key=lambda item: item["rank"])[0]
        # 价目信息来自绑定库，且不依赖跨库外键
        assert primary["record"]["manufacturer"] == "绑定制造商"
        assert primary["record"]["brand"] == "绑定品牌"

        # 客户版只对已确认行写价，先确认再导出
        confirm = client.post(f"/api/quote-lines/{line['id']}/confirm", json={"override_reason": "绑定测试"})
        assert confirm.status_code == 200, confirm.text

        export = client.post(f"/api/quote-jobs/{job['id']}/export/customer")
        assert export.status_code == 200
        sheet = load_workbook(BytesIO(export.content))["询价单"]
        headers = [cell.value for cell in sheet[1]]
        assert "报价1制造商" in headers
        assert sheet.cell(2, headers.index("报价1制造商") + 1).value == "绑定制造商"
        assert sheet.cell(2, headers.index("报价1品牌") + 1).value == "绑定品牌"


def test_quote_job_without_database_key_still_works():
    with TestClient(app) as client:
        login(client)
        created = client.post(
            "/api/quote-jobs",
            data={"requested_option_count": "3"},
            files={"file": ("无绑定询价.xlsx", inquiry_workbook("电子天平", "100g", "台", 1),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert created.status_code == 202, created.text
        job = created.json()
        assert job["database_key"] == ""
        assert job["database_name"] == ""


def test_quote_job_rejects_unknown_database_key():
    with TestClient(app) as client:
        login(client)
        response = client.post(
            "/api/quote-jobs",
            data={"requested_option_count": "3", "database_key": "不存在的库键"},
            files={"file": ("询价.xlsx", inquiry_workbook("电子天平", "100g", "台", 1),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert response.status_code == 404


# ---- 上传健壮性 ------------------------------------------------------------

def test_import_survives_temp_cleanup_failure(monkeypatch):
    """回归：删临时文件失败，绝不能把一次成功的导入变成 500。

    故障现场（2026-09-22 线上实测）：``_with_history_upload`` 把临时文件清理写在
    ``finally`` 里，而受限环境的安全策略会在 ``Path.unlink`` 上抛 ``SystemExit``。
    异常从 ``finally`` 冒泡，FastAPI 直接回 500 —— 用户看到"上传失败"，但崩溃点在
    解析结束、写库之前，数据其实**一条都没进去**，非常容易误判成文件有问题。

    这里直接把"unlink 必炸"注入进去，断言接口仍然 200 且数据真的落库。
    """
    import pathlib

    original_unlink = pathlib.Path.unlink

    def exploding_unlink(self, missing_ok=False):
        if "tmp" in self.parts and self.name.startswith("import-"):
            raise SystemExit(1)
        return original_unlink(self, missing_ok=missing_ok)

    with TestClient(app) as client:
        login(client)
        entry = create_database(client, "清理失败回归库")
        monkeypatch.setattr(pathlib.Path, "unlink", exploding_unlink)
        try:
            response = client.post(
                f"/api/databases/{entry['key']}/import",
                files={"file": ("清理回归.xlsx",
                                price_workbook("清理回归专用量筒", "50ml", "个", 18.0, "回归品牌", "回归制造商"),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
        finally:
            monkeypatch.undo()

        assert response.status_code == 200, response.text
        assert response.json()["inserted"] == 1

        listed = database_keys(client)
        assert listed["清理失败回归库"]["history_count"] == 1, "数据必须真的落库，不能只是接口不报错"


# ---- 统计缓存新鲜度 --------------------------------------------------------

def test_stats_are_not_stale_after_wal_only_write():
    """回归：库列表的条数/大小不能因为"写入只落进 WAL"而长期停在旧值。

    实测故障（2026-09-22）：系统库里明明已经有 1 个报价任务，界面上的库卡片却
    一直显示"任务 0"，库文件大小也是旧值，**直到进程重启才纠正**。

    根因：统计缓存只拿主库 mtime 当键。SQLite 开着 WAL 时写入落进 ``-wal``，
    主库文件本身不被改动、mtime 不变 → 缓存永远命中。而"新建报价任务"这条写路径
    没有（也不该依赖）手动 invalidate_stats，于是计数一直错着。

    修法：缓存键改成"主库 + -wal"的指纹，任何写入都会自然失效。

    为什么用**系统库**来测：报价任务行是经 ``Depends(get_db)`` 写进业务库（= 系统库）的，
    ``database_key`` 只决定"从哪个价目库读价目"，并不会写那个价目库的文件。
    拿一个自建价目库来测，它的文件自始至终没被动过，测出来的是假阴性。
    """
    system_key = system_db_key()
    path = registry.db_path(system_key)

    with TestClient(app) as client:
        login(client)

        # 先读一次把缓存填上，并记下基线与指纹
        baseline = registry.db_stats(system_key)
        fingerprint_before = registry._stats_fingerprint(path)

        created = client.post(
            "/api/quote-jobs",
            data={"requested_option_count": "3"},
            files={"file": ("缓存新鲜度询价.xlsx", inquiry_workbook("缓存新鲜度专用天平", "200g", "台", 1),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert created.status_code == 202, created.text

        # 机制检查：指纹必须能感知到这次写入，否则下面的行为断言没有意义
        assert registry._stats_fingerprint(path) != fingerprint_before, "指纹没变，说明缓存键仍然漏掉 WAL 写入"

        # 行为检查：不调用 invalidate_stats，直接再读一次统计，必须看到新任务
        after = registry.db_stats(system_key)
        assert after["job_count"] == baseline["job_count"] + 1, (
            f"任务数停在旧值 {after['job_count']}（应为 {baseline['job_count'] + 1}），缓存没失效"
        )
